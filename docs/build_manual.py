#!/usr/bin/env python3
"""Build the illustrated TrackRS player and developer manual."""

from __future__ import annotations

import re
import subprocess
import xml.etree.ElementTree as ET
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "TrackRS_User_and_Developer_Manual.docx"
ORANGE = "E65B2F"
INK = "202833"
MUTED = "5B6875"
PALE = "F1F4F6"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def set_keep_with_next(paragraph) -> None:
    paragraph.paragraph_format.keep_with_next = True


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in (begin, instruction, separate, placeholder, end):
        run._r.append(element)


def set_repeat_row_keep(row) -> None:
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, headers, rows, widths=None, font_size=8.5):
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    header = table.rows[0]
    set_repeat_table_header(header)
    for index, label in enumerate(headers):
        cell = header.cells[index]
        set_cell_shading(cell, INK)
        set_cell_margins(cell)
        paragraph = cell.paragraphs[0]
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(str(label))
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor.from_string(WHITE)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        for index, value in enumerate(values):
            cell = cells[index]
            set_cell_margins(cell)
            if row_index % 2:
                set_cell_shading(cell, PALE)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            run.font.size = Pt(font_size)
            run.font.color.rgb = RGBColor.from_string(INK)
        set_repeat_row_keep(table.rows[-1])
    if widths:
        for row in table.rows:
            for cell, width in zip(row.cells, widths):
                cell.width = Inches(width)
    document.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_body(document, text, style=None):
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.add_run(text)
    return paragraph


def add_bullet(document, text, level=0):
    style = "List Bullet" if level == 0 else "List Bullet 2"
    paragraph = document.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)
    return paragraph


def add_number(document, text):
    paragraph = document.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.add_run(text)
    return paragraph


def add_heading(document, text, level=1):
    paragraph = document.add_heading(text, level=level)
    set_keep_with_next(paragraph)
    return paragraph


def add_callout(document, title, text, fill="FFF1E9"):
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=150, start=180, bottom=150, end=180)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title)
    run.bold = True
    run.font.color.rgb = RGBColor.from_string(ORANGE)
    paragraph = cell.add_paragraph(text)
    paragraph.paragraph_format.space_after = Pt(0)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def add_picture(document, relative_path, width=6.3, caption=None):
    path = ROOT / relative_path
    if not path.exists():
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(path), width=Inches(width))
    if caption:
        caption_paragraph = document.add_paragraph()
        caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        caption_paragraph.paragraph_format.space_after = Pt(8)
        run = caption_paragraph.add_run(caption)
        run.italic = True
        run.font.size = Pt(8)
        run.font.color.rgb = RGBColor.from_string(MUTED)


def parse_xml(path):
    return ET.parse(path).getroot()


def read_inventory():
    maps = []
    for path in sorted((ROOT / "data/maps").rglob("*.level")):
        item = parse_xml(path)
        maps.append({
            "name": item.get("name", path.stem),
            "author": item.get("author", "Not stated"),
            "description": item.get("description", "No route description in level file."),
            "path": path.relative_to(ROOT).as_posix(),
        })

    events = []
    stages = []
    referenced_stages = set()
    for path in sorted((ROOT / "data/events").glob("*/*.event")):
        event = parse_xml(path)
        event_stages = []
        for entry in event.findall(".//level"):
            relative = entry.get("file", "")
            stage_path = (path.parent / relative).resolve()
            if not relative or not stage_path.exists():
                continue
            referenced_stages.add(stage_path)
            stage = parse_xml(stage_path)
            stage_info = {
                "event": event.get("name", path.parent.name),
                "name": stage.get("name", stage_path.parent.name),
                "author": stage.get("author", event.get("author", "Not stated")),
                "description": stage.get("description", "Route profile is defined by its terrain and checkpoints."),
                "path": stage_path.relative_to(ROOT).as_posix(),
            }
            event_stages.append(stage_info)
            stages.append(stage_info)
        events.append({
            "name": event.get("name", path.parent.name),
            "author": event.get("author", "Not stated"),
            "locked": event.get("locked", "yes").lower() == "yes",
            "stages": event_stages,
        })

    all_stage_paths = {path.resolve() for path in (ROOT / "data/events").rglob("*.level")}
    orphan_stages = sorted(all_stage_paths - referenced_stages)

    tracked = subprocess.run(
        ["git", "ls-files", "data/vehicles/*/*.vehicle"],
        cwd=ROOT,
        check=True,
        text=True,
        capture_output=True,
    ).stdout.splitlines()
    vehicles = []
    for relative in tracked:
        path = ROOT / relative
        if not path.exists():
            continue
        item = parse_xml(path)
        params = item.find("genparams")
        wheels = item.findall(".//wheel")
        driven = [wheel for wheel in wheels if float(wheel.get("drive", "0")) > 0]
        axles = {
            "front" if float(wheel.get("pos", "0,0,0").split(",")[1]) > 0 else "rear"
            for wheel in driven
        }
        drive = "AWD" if len(axles) > 1 else (f"{next(iter(axles)).upper()} drive" if axles else "not specified")
        mass = float(params.get("mass", "0")) if params is not None else 0
        vehicles.append({
            "folder": path.parent.name,
            "name": item.get("name", path.parent.name),
            "class": item.get("class", "Unspecified"),
            "mass": mass,
            "drive": drive,
            "wheel_count": len(driven),
            "path": relative,
        })

    plugin_paths = sorted((ROOT / "data/plugins").glob("*.zip"))
    credits = (ROOT / "doc/DATA_AUTHORS.txt").read_text(errors="replace")
    unlicensed = set(re.findall(r"plugins/(\S+\.zip) \(unlicen[cs]ed\)", credits, re.IGNORECASE))
    plugins = [{
        "name": path.stem,
        "status": "Marked unlicensed in credits" if path.name in unlicensed else "Check credits and permission before redistribution",
    } for path in plugin_paths]

    ahead = parse_xml(ROOT / "data/maps/ahead/ahead.level")
    vehicle_options = [node.get("type", "") for node in ahead.findall(".//vehicleoption")]
    existing_vehicle_files = {item["path"].removeprefix("data/") for item in vehicles}
    missing_options = [option for option in vehicle_options if option.lstrip("/") not in existing_vehicle_files]

    return maps, events, stages, orphan_stages, vehicles, plugins, missing_options


def route_advice(description: str) -> str:
    text = description.lower()
    advice = []
    if any(word in text for word in ("narrow", "tight", "curvy", "hairpin", "chicane")):
        advice.append("Brake before the turn; use short, precise steering inputs.")
    if any(word in text for word in ("fast", "speedy", "wide")):
        advice.append("Build speed only where the next bend is visible; lift early for blind corners.")
    if any(word in text for word in ("jump", "jumpy")):
        advice.append("Straighten the car before crests and avoid abrupt steering while airborne.")
    if any(word in text for word in ("ice", "snow", "slippery")):
        advice.append("Use gentle throttle and steering; leave extra room to slow down.")
    if any(word in text for word in ("fog", "visibility", "low light", "dark")):
        advice.append("Read pace notes, reduce speed, and look farther ahead than the near road edge.")
    if any(word in text for word in ("sand", "mud", "wet", "damp")):
        advice.append("Keep momentum smooth; avoid sudden throttle or braking that unsettles the car.")
    if any(word in text for word in ("gravel", "dirt", "earth", "offroad")):
        advice.append("Expect a loose surface; steer progressively and avoid over-correcting slides.")
    return " ".join(dict.fromkeys(advice)) or "Scout the route, follow pace notes, and keep inputs smooth over changing terrain."


def vehicle_character(vehicle):
    weight = vehicle["mass"]
    car_class = vehicle["class"].lower()
    family_key = vehicle["folder"].split("_", 1)[0]
    family = {
        "cordo": "Cordo's documented tune is the middle ground; use it as the balanced reference when learning a stage.",
        "evo": "The project's tuning note describes Evo as circuit-oriented; favor clean lines, early braking, and a settled exit.",
        "fox": "The project's tuning note describes Fox as offroad-oriented; use its loose-surface character on dirt and gravel, but still slow down for ice and tight corners.",
    }.get(family_key, "Use smooth inputs and tune your pace to the stage surface.")
    if "kit" in car_class:
        profile = "Light class; quick to change direction."
    elif "s500" in car_class or "super" in car_class:
        profile = "Heavy class; prioritize early braking and settled corner entry."
    else:
        profile = "Rally class; balance speed with stable, early turn-in."
    drive = vehicle["drive"]
    if drive == "AWD":
        approach = "Four driven wheels in the file: progressive throttle can help pull the car out of a corner."
    elif "REAR" in drive:
        approach = "Rear-wheel drive in the file: ease on power mid-corner to avoid rotating the rear."
    elif "FRONT" in drive:
        approach = "Front-wheel drive in the file: avoid asking for full throttle and maximum steering at once."
    else:
        approach = "Drive layout is not identified in the vehicle file; use measured inputs."
    return f"{family} {profile} {approach}"


def configure_document(document):
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    section.header_distance = Inches(0.3)
    section.footer_distance = Inches(0.3)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size, color in (("Title", 34, INK), ("Heading 1", 22, ORANGE), ("Heading 2", 14, INK), ("Heading 3", 11, ORANGE)):
        style = document.styles[style_name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = header.add_run("TRACKRS  |  PLAYER + DEVELOPER MANUAL")
    run.bold = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("TrackRS 0.6.7  |  Workspace manual  |  ")
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_page_number(footer)


def add_code_block(document, text):
    text = text.replace("\n+", "\n")
    table = document.add_table(rows=1, cols=1)
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F3F5F7")
    set_cell_margins(cell, top=110, start=140, bottom=110, end=140)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.font.name = "Consolas"
    run.font.size = Pt(8)
    document.add_paragraph().paragraph_format.space_after = Pt(1)


def build_manual():
    maps, events, stages, orphan_stages, vehicles, plugins, missing_options = read_inventory()
    document = Document()
    configure_document(document)

    # Cover
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = ROOT / "data/icon/logo.png"
    if logo.exists():
        paragraph.add_run().add_picture(str(logo), width=Inches(2.6))
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(16)
    run = paragraph.add_run("TRACKRS")
    run.bold = True
    run.font.size = Pt(34)
    run.font.color.rgb = RGBColor.from_string(INK)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("PLAYER'S GUIDE  /  DEVELOPER'S HANDBOOK")
    run.bold = True
    run.font.size = Pt(15)
    run.font.color.rgb = RGBColor.from_string(ORANGE)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.add_run("Rally driving, vehicles, maps, installation, source, and debugging\n")
    run = paragraph.add_run("Version 0.6.7  |  Prepared 24 October 2026")
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor.from_string(MUTED)
    add_picture(document, "data/maps/ahead/ss.png", 6.25, "Ahead: the long gravel-and-hill stage used by the demo.")
    add_callout(document, "Read this first", "This guide describes the repository snapshot used to build it. The demo menu unlocks only Ahead; some vehicle references in Ahead do not resolve in this checkout. Full data folders are not the same as content accessible in the demo.")
    document.add_page_break()

    # Contents
    add_heading(document, "In this manual", 1)
    contents = [
        "1. What TrackRS is and who made it",
        "2. Install, launch, and first race",
        "3. Controls and rally driving technique",
        "4. Demo access, maps, events, and add-on packs",
        "5. Cars and how their handling differs",
        "6. Settings, profiles, and troubleshooting",
        "7. Development, building, debugging, and content creation",
        "8. Credits, licenses, and distribution notes",
        "Appendix A. All 132 standalone map files",
        "Appendix B. Event-stage inventory and optional packs",
    ]
    for item in contents:
        add_bullet(document, item)
    add_callout(document, "Scope note", "TrackRS is a rebranded and modernized fork of Trigger Rally. The creator identified in the repository is Saad AIT YAHIA (@Saad-programmer). The map and vehicle inventories below are read from this checkout; availability can differ by build and by the rights attached to individual data files.", "EAF3F6")

    # 1
    add_heading(document, "1. Meet TrackRS", 1)
    add_body(document, "TrackRS is a single-player 3D rally game built in C++11. Race point-to-point stages against target times, follow spoken and visual co-driver pace notes, and manage grip across gravel, dirt, tarmac, sand, snow, ice, mud, and water. Terrain, weather, fog, lighting, and visibility vary by level.")
    add_body(document, "The game descends from Trigger Rally and retains its map/event format and many community-created assets. The TrackRS rebrand adds its application identity and menu presentation. The documented release is version 0.6.7; older version-history entries in the repository describe development through the 0.6.6.x era.")
    add_heading(document, "What a race asks you to do", 2)
    for text in [
        "Choose a level and read its route description, author, minimap, and screenshot.",
        "Choose one of the vehicles declared by that level, when valid choices are available.",
        "Start from the countdown, drive through checkpoints in order, and finish before the target time.",
        "Use the co-driver's notes to prepare for turns hidden by crests, fog, or terrain.",
        "Save best times to the local player profile; optionally record and replay a ghost car.",
    ]:
        add_number(document, text)
    add_picture(document, "data/events/01-triggercup/far/ss.png", 6.1, "A Trigger Cup stage screenshot bundled with the game data.")

    # 2
    add_heading(document, "2. Install and start", 1)
    add_heading(document, "Linux from source", 2)
    add_body(document, "The game can run from the build tree; a system installation is not required. Install the compiler and development packages first.")
    add_code_block(document, "sudo apt install build-essential libsdl2-dev libsdl2-image-dev libglew-dev \\\n+  libopenal-dev libalut-dev libtinyxml2-dev libphysfs-dev \\\n+  libgl1-mesa-dev libglu1-mesa-dev\n+cd src\n+make\n+cd ../bin\n+./trackrs")
    add_heading(document, "Debian demo package", 2)
    add_code_block(document, "sudo apt install ./trackrs-demo_0.6.7_amd64.deb")
    add_body(document, "The package installs the executable and default configuration under /usr/local/games, game data under /usr/local/share/games/trackrs, and a desktop entry. Package dependencies must be available for your distribution. See the package's installed copyright and data-credit documents before redistribution.")
    add_heading(document, "Windows", 2)
    add_body(document, "The project documents MSYS2 makefiles for 32-bit and 64-bit builds. The NSIS installer expects a Windows executable and its required DLLs in bin/ before installer.nsi is compiled with makensis. Refer to doc/README-win.txt and doc/BUILDING.txt; Visual Studio is not the documented toolchain.")
    add_heading(document, "First launch", 2)
    for text in [
        "Start TrackRS and choose Single Race. In this demo, select Ahead.",
        "Read the route description and inspect the minimap before confirming the race.",
        "Select a vehicle if the list is populated, then wait for the countdown.",
        "Drive through checkpoints in order. The HUD shows time, checkpoint progress, minimap, and speed/gear.",
        "Finish under the target time. Use Recover if the car is stranded; use Recover at checkpoint to restart from the previous checkpoint.",
    ]:
        add_number(document, text)
    if missing_options:
        add_callout(document, "Known vehicle-reference issue in this checkout", "Ahead currently declares these vehicle paths, but corresponding .vehicle files are absent: " + ", ".join(missing_options) + ". The repository's tracked roster uses other paths. If the car chooser is empty or the stage cannot start, this is a data/configuration mismatch to fix in data/maps/ahead/ahead.level before testing the build.")

    # 3
    add_heading(document, "3. Controls and driving", 1)
    add_picture(document, "Screenshot from 2026-10-24 15-35-27.png", 6.2, "In-race HUD example: timer, co-driver note, checkpoint count, minimap, gear, and speedometer.")
    add_table(document, ["Key", "Action", "Use"], [
        ("Up", "Accelerate", "Hold for power; on keyboard, tap or release to meter throttle on loose surfaces."),
        ("Down", "Foot brake / reverse", "Brake before turns. Continue holding after stopping to reverse."),
        ("Left / Right", "Steer", "Use short corrections; avoid sawing the wheel at speed."),
        ("Space", "Handbrake", "Useful for a tight low-speed rotation; release promptly to regain traction."),
        ("R", "Recover car", "Return the car to the start/recovery position when stuck."),
        ("Q", "Recover at checkpoint", "Return to the previous checkpoint after a mistake."),
        ("C", "Change camera", "Cycle chase, bumper, side, hood, and periscope views."),
        (", / .", "Rotate camera", "Adjust the third-person camera."),
        ("M", "Toggle minimap", "Show or hide route guidance."),
        ("P", "Pause race", "Pause or resume the current race."),
        ("N", "Toggle interface", "Hide or restore the HUD."),
        ("K", "Toggle checkpoints", "Show or hide checkpoint indicators."),
        ("F12", "Save screenshot", "Save a screenshot in the TrackRS user folder."),
    ], [1.0, 1.55, 4.1], 8.2)
    add_heading(document, "Reading pace notes", 2)
    add_body(document, "Notes use direction and severity words such as left, right, easy, medium, hard, flat, square, hairpin, chicane, long, cut, don't cut, jump, over, and into. Read them as preparation, not as a command to turn at a fixed instant: brake before a tight or blind corner, then turn as the road arrives. Pair the spoken call with the icon and minimap if enabled.")
    add_heading(document, "A repeatable corner routine", 2)
    for text in [
        "Look ahead and identify the next pace note, bend, crest, or surface change.",
        "Lift off the accelerator early; brake while the car is straight and stable.",
        "Turn progressively toward the inside of the bend. Avoid a large late steering correction.",
        "As the car points toward the exit, unwind steering and feed in throttle smoothly.",
        "On gravel, snow, sand, mud, or wet ground, leave more stopping distance and accept lower corner speed.",
        "After a missed checkpoint or rollover, recover and continue; use Q if restarting from the last checkpoint is safer.",
    ]:
        add_number(document, text)
    add_callout(document, "Keyboard vs. analog input", "Keyboard controls are digital: a key is either pressed or released. Smooth throttle is easier with a wheel or controller axis. Joystick examples are present in the default config but disabled; enable and calibrate them before racing.", "EAF3F6")
    add_picture(document, "data/maps/authentic/ss.png", 6.0, "A narrow, jump-prone route rewards early braking and clean steering.")

    # 4
    add_heading(document, "4. Demo access and the world of TrackRS", 1)
    add_heading(document, "What the demo unlocks", 2)
    add_body(document, "The current menu code disables Events and Practice and marks every standalone level except Ahead as locked. Selection and start actions are guarded in src/Trigger/menu.cpp, so merely copying map files into data/ does not unlock them. The full data tree is much larger than the playable demo.")
    add_body(document, "The source tree contains " + str(len(maps)) + " standalone map definition(s), " + str(len(events)) + " event definition(s), and " + str(len(stages)) + " event-stage files referenced by event definitions. These are inventory counts, not a promise that the demo menu can launch them.")
    add_heading(document, "The demo stage: Ahead", 2)
    ahead = next((item for item in maps if item["name"].lower() == "ahead"), None)
    if ahead:
        add_body(document, "Ahead is credited to " + ahead["author"] + ". The level describes it as " + ahead["description"] + ". Its target time is 1,500 seconds (25 minutes), and the file declares " + str(len(parse_xml(ROOT / "data/maps/ahead/ahead.level").findall(".//checkpoint"))) + " checkpoints. Study the minimap and co-driver notes; it is a long route, so avoid unnecessary recoveries.")
    add_heading(document, "Events", 2)
    event_rows = []
    for event in events:
        state = "Locked in demo" if event["locked"] or event["name"].startswith("01.") else "Locked in demo"
        event_rows.append((event["name"], event["author"], len(event["stages"]), state))
    add_table(document, ["Event", "Credited author", "Referenced stages", "Demo status"], event_rows, [2.0, 1.75, 1.2, 1.7], 8.0)
    add_callout(document, "Unlocking content for development", "To test the full map/event flow, work in a developer build and review the demo-only menu guards in src/Trigger/menu.cpp. Restore normal level selection and event/practice navigation, then test event locks, level vehicle options, saves, and source/data licensing. This is a source modification, not a user-facing in-game unlock code.")
    if orphan_stages:
        add_body(document, "Inventory note: " + str(len(orphan_stages)) + " event-stage level file(s) are present but not referenced by an event file: " + ", ".join(path.relative_to(ROOT).as_posix() for path in orphan_stages) + ".")

    # 5
    add_heading(document, "5. Cars and their handling", 1)
    add_body(document, "The tracked release tree contains nine vehicle definitions: three families across Kit-Car, S500 (Super 500), and WRC classes. The figures below come from each .vehicle file. Mass and driven-wheel layout describe the data; handling advice is a practical interpretation, not a benchmark or guarantee of lap time.")
    rows = []
    for vehicle in vehicles:
        rows.append((vehicle["name"], vehicle["class"], f'{vehicle["mass"]:,.0f} kg', vehicle["drive"], vehicle_character(vehicle)))
    add_table(document, ["Car identity", "Class", "Mass", "Driven wheels", "Character and driving approach"], rows, [1.15, 0.8, 0.65, 0.85, 3.2], 7.7)
    add_heading(document, "How to choose", 2)
    add_bullet(document, "Kit-Car: the three tracked models are the lightest entries in their families. Start here when learning corner timing and pace notes.")
    add_bullet(document, "S500: all three tracked definitions are heavier and drive all four wheels. Brake earlier; use measured throttle to exploit traction without pushing wide.")
    add_bullet(document, "WRC: the tracked definitions vary substantially in mass and are rear-wheel drive in their XML. Keep the car settled before power, especially on slippery exits.")
    add_bullet(document, "The original project notes summarize the tuning identities as 'Fox offroad, Evo circuit, Cordo in between.' Treat that as a broad setup flavor, not a substitute for checking each level's allowed-car list.")
    add_callout(document, "Only valid choices can appear", "A level's <vehicleoption> entries control its car list. The nine tracked cars listed here are data assets; Ahead's current options point to three missing definitions. Also present in this local checkout are untracked folders evo_rally and lancer_evo. They are not part of the tracked nine-car roster and are not represented as official release vehicles in this manual.", "EAF3F6")
    add_picture(document, "data/maps/lazyhills/ss.png", 6.0, "A game-map screenshot; the actual car model is assembled from its .vehicle and .obj assets.")

    # 6
    add_heading(document, "6. Settings, profiles, and troubleshooting", 1)
    add_heading(document, "Player profile and saves", 2)
    add_body(document, "On Linux, the config and player data live under ~/.trackrs/; the default configuration names the profile trackrs-0.6.x.config. Player names should use letters, digits, underscore, and spaces. Best times, player profiles, and screenshots are stored in this user area. Deleting the config resets options; back up profiles and records before removing the entire folder.")
    add_heading(document, "Useful defaults", 2)
    add_table(document, ["Setting", "Default", "Practical note"], [
        ("Speed unit", "mph", "Change to kph in the parameters section."),
        ("Driving assist", "0.33", "Tune assistance if steering is difficult; the value is a game parameter, not a difficulty selector."),
        ("Engine / effects / co-driver volume", "0.2 / 0.7 / 1.0", "Each volume is in the 0-to-1 range."),
        ("Weather / foliage / dirt effects", "On", "Disable costly effects on slower graphics hardware."),
        ("Ghost recording", "Off", "Enable to record and replay a run."),
        ("Joystick templates", "Disabled", "Enable one mapping and adjust axis deadzones for your device."),
    ], [2.1, 1.25, 4.0], 8.2)
    add_heading(document, "Troubleshooting", 2)
    add_table(document, ["Symptom", "Checks"], [
        ("Game cannot find data", "Keep trackrs.config.defs beside the executable and confirm its <datadirectory> points to the installed data folder."),
        ("Empty car list / level fails to start", "Check every <vehicleoption> path in the .level file against an existing data/vehicles/<folder>/*.vehicle file. Ahead currently has three unresolved references in this checkout."),
        ("Profile not loading", "Use the supported player-name characters; inspect the profile XML and preserve a backup."),
        ("Low frame rate", "Disable dirteffect, weather, or foliage; reduce anisotropy; use point snowflakes."),
        ("Snow artifacting", "Set snowflaketype to square."),
        ("No intro video", "The optional video needs ffplay; Escape or Q skips playback."),
        ("Uncomfortable stereo", "Turn stereo off or reduce stereoeyeseparation. Do not use stereo if it causes discomfort or conflicts with a health condition."),
    ], [2.0, 5.35], 8.0)

    # 7
    add_heading(document, "7. Development, building, and debugging", 1)
    add_heading(document, "Code architecture", 2)
    add_table(document, ["Layer", "Responsibilities", "Useful files"], [
        ("PEngine", "SDL application loop, rendering, terrain, textures, models, audio, config, PhysFS, math.", "src/PEngine/app.cpp, render.cpp, terrain.cpp, audio.cpp, config.cpp"),
        ("PSim", "Rigid bodies, fixed-step world, vehicle/wheel/suspension behavior, engines, collision, damage.", "src/PSim/sim.cpp, rigidbody.cpp, vehicle.cpp, engine.cpp, collision.cpp"),
        ("Trigger", "Menus, event/level flow, controls, timer/checkpoints, camera/HUD, ghost recording.", "src/Trigger/main.cpp, game.cpp, menu.cpp, control.cpp, ghost.cpp"),
    ], [1.0, 3.7, 2.65], 8.0)
    add_body(document, "The application uses C++11 and GNU Make. Core libraries include SDL2, SDL2_image, OpenGL/GLEW, OpenAL/ALUT, PhysicsFS, and TinyXML-2. The project uses XML for configuration, levels, events, and vehicle definitions; image maps and Wavefront OBJ/MTL files provide much of the visual content.")
    add_heading(document, "Linux developer build", 2)
    add_code_block(document, "cd src\nmake\ncd ../bin\n./trackrs")
    add_body(document, "The makefile enables -Wall -Wextra -pedantic and defaults to an optimized native build. For a local debug build, first clean existing objects, then override optimization flags:")
    add_code_block(document, "cd src\nmake clean\nOPTIMS='-O0 -g' make\ncd ../bin\ngdb ./trackrs\n# In GDB: run\n# After a crash: bt")
    add_body(document, "The source build documentation describes INDEVEL as enabling additional terrain and co-driver diagnostics. Build with OPTIMS='-O0 -g -DINDEVEL' when you need those diagnostics, and compare results against a normal release build.")
    add_heading(document, "Debugging checklist", 2)
    for text in [
        "Reproduce with the exact level, vehicle definition, config, and input device; record the route and steps.",
        "Check the game log/output for missing XML, model, texture, sound, or vehicle paths.",
        "For a crash, reproduce in a debug build and capture a GDB backtrace; include symbols and the exact build flags.",
        "For physics differences, compare one variable at a time: mass, power curve, gear ratios, wheel drive/steer/brake, suspension, friction, and terrain type.",
        "For a map bug, inspect the level XML and its height, color, foliage, road, terrain, minimap, and screenshot images.",
        "For missing cars, verify every level vehicleoption resolves to a parseable .vehicle and every referenced OBJ/MTL/texture file exists.",
        "Keep source changes separate from data changes; document the file, date, author, and license of every new or modified content asset.",
    ]:
        add_bullet(document, text)
    add_heading(document, "Creating maps and events", 2)
    add_body(document, "A .level describes terrain maps, weather, foliage, vehicle options, start pose, target time, checkpoints, and co-driver notes. Heightmaps should be PNG to avoid compression artifacts changing the terrain. The .event file lists levels, author, initial lock state, and the next event to unlock. Copy a similar existing level/event, use unique paths, test each asset, and add authorship/license details to doc/DATA_AUTHORS.txt.")
    add_heading(document, "Creating vehicles", 2)
    add_body(document, "Vehicle definitions connect a name/class and mass to a drive system, engine curve, gearbox, body model, wheel models, suspension, brakes, steering, and collision clips. TrackRS reads OBJ meshes with the documented restrictions: triangles, one material per mesh, body geometry together, and wheels in separate files. Start by duplicating a known-good vehicle and change one parameter at a time.")
    add_code_block(document, "# Inspect a level's declared vehicles and checkpoints\npython3 - <<'PY'\nimport xml.etree.ElementTree as ET\nroot = ET.parse('data/maps/ahead/ahead.level').getroot()\nfor item in root.findall('.//vehicleoption'):\n    print(item.get('type'))\nprint('checkpoints:', len(root.findall('.//checkpoint')))\nPY")
    add_heading(document, "Packaging", 2)
    add_body(document, "The repository supplies build_deb.sh for the Debian demo and installer.nsi for NSIS on Windows. Build from the repository root after building the Linux executable. Verify package contents, dependencies, installed data paths, licenses, and third-party credits. Store catalogs may require separate repository metadata and hosted screenshots; packaging screenshots alone does not publish a store listing.")
    add_callout(document, "Working-copy changes", "This manual was generated from a working tree that contains local uncommitted edits in src/PSim/vehicle.cpp and src/Trigger/game.cpp related to rejecting invalid vehicle models and skipping failed vehicle loads. Those edits are not represented as an upstream release guarantee; compare the source tree you build before relying on this behavior.", "FFF7E6")

    # 8
    add_heading(document, "8. Credits, licenses, and safe redistribution", 1)
    add_body(document, "TrackRS identifies Saad AIT YAHIA as creator and maintainer. It is based on Trigger Rally and includes contributions from the upstream community, including map authors Onsemeliot, Andrei Bondor, Richard Langridge, and the many additional contributors listed in doc/DATA_AUTHORS.txt.")
    add_body(document, "The project states that its code is licensed under GNU GPL version 2 or, at the recipient's option, a later version. The complete text is doc/COPYING.txt. GPL obligations include preserving notices and supplying the license; when distributing modified covered files, GPLv2 section 2(a) requires prominent notices that the files were changed and the date of the change.")
    add_callout(document, "Data is not one license", "Game data includes assets with different authors and licenses. The project credits explicitly mark some optional plugin archives as unlicensed. Do not redistribute or publish those works unless you have confirmed permission. A GPL license for code does not automatically relicense separate art, audio, maps, fonts, or third-party libraries.")
    add_body(document, "Before publishing a store build, audit every shipped file, confirm the rights and attribution requirements for code, maps, car artwork, audio, fonts, and screenshots, provide corresponding source where the GPL requires it, and ask qualified counsel or the relevant rights holders about unresolved claims. This manual is informational, not legal advice.")
    add_body(document, "Key project references: README.md, doc/README.txt, doc/BUILDING.txt, doc/README-stereo.txt, doc/COPYING.txt, doc/DATA_AUTHORS.txt, data/metainfo/, src/GNUmakefile, and src/Trigger/menu.cpp.")

    # Appendix A: full map list
    document.add_page_break()
    add_heading(document, "Appendix A. Standalone map inventory", 1)
    add_body(document, "The following map names, route descriptions, and authors are read from data/maps/*.level in this checkout. All but Ahead are locked by the current demo menu. Suggested approaches are derived from the route descriptions; inspect each level in-game when using a full-content developer build.")
    map_rows = []
    for item in maps:
        description = item["description"].strip() or "No route description provided."
        if item["name"].lower() == "ahead":
            status = "DEMO MAP"
        else:
            status = "Full-data map; locked in demo"
        profile = f'{description}\nApproach: {route_advice(description)}\nAuthor: {item["author"]} | {status}'
        map_rows.append((item["name"], profile))
    add_table(document, ["Map", "Route profile, suggested approach, author, access"], map_rows, [1.45, 5.9], 7.6)

    # Appendix B: stages and plugins
    add_heading(document, "Appendix B. Event stages and optional packs", 1)
    add_body(document, "Event-stage files are listed below when referenced by a .event definition. Stage names and route descriptions come from their .level files. The demo does not expose these events through the menu.")
    stage_rows = [(stage["event"], stage["name"], f'{stage["description"]} | {stage["author"]}') for stage in stages]
    add_table(document, ["Event", "Stage", "Route profile / author"], stage_rows, [1.75, 1.4, 4.2], 7.3)
    if orphan_stages:
        add_body(document, "Unreferenced event-stage file(s): " + ", ".join(path.relative_to(ROOT).as_posix() for path in orphan_stages) + ".")
    add_heading(document, "Optional plugin archives", 2)
    add_body(document, "These ZIP files are present under data/plugins. Their presence does not mean the demo activates them or that redistribution rights are cleared. Check doc/DATA_AUTHORS.txt for each pack before use or redistribution.")
    plugin_rows = [(plugin["name"], plugin["status"]) for plugin in plugins]
    add_table(document, ["Pack", "Credit/license review"], plugin_rows, [2.1, 5.25], 7.8)

    document.core_properties.title = "TrackRS Player and Developer Manual"
    document.core_properties.subject = "TrackRS 0.6.7 demo, vehicles, maps, development, and debugging"
    document.core_properties.author = "TrackRS project documentation"
    document.core_properties.keywords = "TrackRS, Trigger Rally, rally game, manual, development"
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")
    print(f"Inventory: {len(maps)} standalone maps, {len(events)} events, {len(stages)} referenced stages, {len(vehicles)} tracked vehicles, {len(plugins)} plugin archives")


if __name__ == "__main__":
    build_manual()